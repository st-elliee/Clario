from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from typing import Optional
from database import get_connection
from models import *
from auth import verify_password, hash_password, create_access_token, decode_token
from encryption import encrypt, decrypt, hash_afm
from contextlib import asynccontextmanager
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from fastapi import FastAPI, HTTPException, Depends, status, UploadFile, File, Form, Header
from fastapi.staticfiles import StaticFiles
import os
import uuid
import shutil
import json
import secrets

# Demo mode (live demo): file uploads are disabled and the demo accounts
# cannot be renamed, re-passworded or deactivated.
DEMO_MODE = os.getenv("DEMO_MODE", "false").lower() == "true"
DEMO_PROTECTED_USERNAMES = {"admin", "employee"}


def update_overdue_services():
    conn = get_connection()
    if not conn:
        return
    try:
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE client_service
            SET status = 'overdue'
            WHERE status IN ('pending', 'in_progress')
            AND due_date IS NOT NULL
            AND due_date < CURDATE()
            AND service_id != 1
        """)
        conn.commit()
    finally:
        cursor.close()
        conn.close()

@asynccontextmanager
async def lifespan(app: FastAPI):
    update_overdue_services()
    yield

app = FastAPI(lifespan=lifespan)

UPLOAD_DIR = os.getenv("UPLOAD_DIR", os.path.join(os.path.dirname(__file__), "uploads"))
if not DEMO_MODE:
    # In the live demo (serverless, read-only filesystem) uploads are disabled.
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")

CORS_ORIGINS = [
    o.strip()
    for o in os.getenv("CORS_ORIGINS", "http://localhost:5173").split(",")
    if o.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

security = HTTPBearer()

def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    token = credentials.credentials
    payload = decode_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Μη έγκυρο token")
    return payload

def get_admin_user(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Απαιτούνται δικαιώματα admin")
    return current_user


# ============================================================
# ΚΑΤΑΓΡΑΦΗ ΚΙΝΗΣΕΩΝ
# ============================================================
def log_action(employee_id: int, action_type: str, client_id: int = None, details: dict = None):
    try:
        conn = get_connection()
        if not conn:
            return
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO action (employee_id, action_type, client_id, details)
            VALUES (%s, %s, %s, %s)
        """, (
            employee_id,
            action_type,
            client_id,
            json.dumps(details, ensure_ascii=False) if details else None
        ))
        conn.commit()
        cursor.close()
        conn.close()
    except Exception as e:
        print(f"Log error: {e}")


# ============================================================
# AUTH
# ============================================================
@app.post("/api/login", response_model=TokenResponse)
def login(request: LoginRequest):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT * FROM employee WHERE username = %s AND is_active = TRUE",
            (request.username,)
        )
        employee = cursor.fetchone()
        if not employee or not verify_password(request.password, employee["password_hash"]):
            raise HTTPException(status_code=401, detail="Λάθος username ή password")
        cursor.execute(
            "UPDATE employee SET last_login = NOW() WHERE employee_id = %s",
            (employee["employee_id"],)
        )
        conn.commit()

        access_types = []
        if employee["role"] == "admin":
            access_types = ["services"]
        else:
            access_types = ["services"]

        token = create_access_token({
            "sub": str(employee["employee_id"]),
            "username": employee["username"],
            "role": employee["role"],
            "full_name": employee["full_name"],
            "access": access_types
        })
        return TokenResponse(
            access_token=token,
            token_type="bearer",
            employee_id=employee["employee_id"],
            full_name=employee["full_name"],
            role=employee["role"],
            access=access_types
        )
    finally:
        cursor.close()
        conn.close()


# ============================================================
# CLIENTS
# ============================================================
@app.get("/api/clients")
def get_clients(
    search: Optional[str] = None,
    client_type: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        type_filter = " AND c.client_type = %s" if client_type else ""
        type_params = [client_type] if client_type else []

        if search:
            cursor.execute(f"""
                SELECT c.client_id, c.last_name, c.first_name, c.company_name,
                       c.client_type, c.phone, c.email, c.doy, c.profession, c.is_active,
                       cs.status AS declaration_status
                FROM client c
                LEFT JOIN client_service cs ON c.client_id = cs.client_id
                    AND cs.service_id = 1
                    AND YEAR(cs.created_at) = YEAR(CURDATE())
                WHERE c.is_active = TRUE
                  AND (c.last_name LIKE %s
                    OR c.first_name LIKE %s
                    OR c.company_name LIKE %s
                    OR CONCAT(c.last_name, ' ', c.first_name) LIKE %s
                    OR CONCAT(c.first_name, ' ', c.last_name) LIKE %s)
                {type_filter}
                ORDER BY
                    CASE WHEN c.last_name LIKE %s THEN 0 ELSE 1 END,
                    c.last_name ASC,
                    c.first_name ASC
            """, (f"%{search}%", f"%{search}%", f"%{search}%", f"%{search}%", f"%{search}%", f"{search}%") + tuple(type_params))
            results = cursor.fetchall()
            if not results and not client_type:
                afm_hash = hash_afm(search)
                cursor.execute("""
                    SELECT c.client_id, c.last_name, c.first_name, c.company_name,
                           c.client_type, c.phone, c.email, c.doy, c.profession, c.is_active,
                           cs.status AS declaration_status
                    FROM client c
                    LEFT JOIN client_service cs ON c.client_id = cs.client_id
                        AND cs.service_id = 1
                        AND YEAR(cs.created_at) = YEAR(CURDATE())
                    WHERE c.is_active = TRUE AND c.afm_hash = %s
                """, (afm_hash,))
                results = cursor.fetchall()
        else:
            cursor.execute(f"""
                SELECT c.client_id, c.last_name, c.first_name, c.company_name,
                       c.client_type, c.phone, c.email, c.doy, c.profession, c.is_active,
                       cs.status AS declaration_status
                FROM client c
                LEFT JOIN client_service cs ON c.client_id = cs.client_id
                    AND cs.service_id = 1
                    AND YEAR(cs.created_at) = YEAR(CURDATE())
                WHERE c.is_active = TRUE
                {type_filter}
                ORDER BY c.last_name ASC
            """, tuple(type_params))
            results = cursor.fetchall()
        return results
    finally:
        cursor.close()
        conn.close()


@app.get("/api/clients/{client_id}")
def get_client(client_id: int, current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT client_id, last_name, first_name, company_name,
                   client_type, phone, email, address, doy,
                   profession, kad, notes, is_active, created_at,
                   afm, amka, id_number, taxis_username, taxis_password
            FROM client
            WHERE client_id = %s AND is_active = TRUE
        """, (client_id,))
        client = cursor.fetchone()
        if not client:
            raise HTTPException(status_code=404, detail="Ο πελάτης δεν βρέθηκε")
        client["afm"] = decrypt(client["afm"]) if client["afm"] else None
        client["amka"] = decrypt(client["amka"]) if client["amka"] else None
        client["id_number"] = decrypt(client["id_number"]) if client["id_number"] else None
        client["taxis_username"] = decrypt(client["taxis_username"]) if client["taxis_username"] else None
        client["taxis_password"] = decrypt(client["taxis_password"]) if client["taxis_password"] else None
        return client
    finally:
        cursor.close()
        conn.close()


@app.post("/api/clients", status_code=201)
def create_client(client: ClientCreate, current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            INSERT INTO client (afm, afm_hash, last_name, first_name, company_name, client_type,
                phone, email, address, amka, id_number, taxis_username, taxis_password,
                doy, profession, kad, notes, created_by)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """, (
            encrypt(client.afm),
            hash_afm(client.afm),
            client.last_name, client.first_name, client.company_name,
            client.client_type, client.phone, client.email, client.address,
            encrypt(client.amka) if client.amka else None,
            encrypt(client.id_number) if client.id_number else None,
            encrypt(client.taxis_username) if client.taxis_username else None,
            encrypt(client.taxis_password) if client.taxis_password else None,
            client.doy, client.profession, client.kad, client.notes,
            int(current_user["sub"])
        ))
        conn.commit()
        new_id = cursor.lastrowid
        cursor.execute("""
            INSERT INTO action (employee_id, action_type, client_id, ip_address)
            VALUES (%s, 'add_client', %s, '127.0.0.1')
        """, (int(current_user["sub"]), new_id))
        conn.commit()
        return {"message": "Ο πελάτης δημιουργήθηκε επιτυχώς", "client_id": new_id}
    finally:
        cursor.close()
        conn.close()


@app.put("/api/clients/{client_id}")
def update_client(
    client_id: int,
    client: ClientUpdate,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        raw = client.model_dump()
        fields = {}
        for k, v in raw.items():
            if v is None:
                continue
            if k == 'afm':
                fields['afm'] = encrypt(v)
                fields['afm_hash'] = hash_afm(v)
            elif k in ('amka', 'taxis_username', 'taxis_password'):
                fields[k] = encrypt(v)
            else:
                fields[k] = v
        if not fields:
            raise HTTPException(status_code=400, detail="Δεν δόθηκαν στοιχεία προς ενημέρωση")
        set_clause = ", ".join([f"{k} = %s" for k in fields.keys()])
        values = list(fields.values()) + [client_id]
        cursor.execute(f"UPDATE client SET {set_clause} WHERE client_id = %s", values)
        conn.commit()
        cursor.execute("""
            INSERT INTO action (employee_id, action_type, client_id, ip_address)
            VALUES (%s, 'edit_client', %s, '127.0.0.1')
        """, (int(current_user["sub"]), client_id))
        conn.commit()
        return {"message": "Ο πελάτης ενημερώθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


@app.delete("/api/clients/{client_id}")
def delete_client(
    client_id: int,
    current_user: dict = Depends(get_current_user)
):
    if current_user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Μόνο ο admin μπορεί να διαγράψει πελάτες")
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor()
        cursor.execute("UPDATE client SET is_active = FALSE WHERE client_id = %s", (client_id,))
        conn.commit()
        cursor.execute("""
            INSERT INTO action (employee_id, action_type, client_id, ip_address, details)
            VALUES (%s, 'delete_client', %s, '127.0.0.1', '{}')
        """, (int(current_user["sub"]), client_id))
        conn.commit()
        return {"message": "Ο πελάτης διαγράφηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


# ============================================================
# CLIENT SERVICES
# ============================================================
@app.get("/api/clients/{client_id}/services")
def get_client_services(client_id: int, current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT cs.client_id, cs.service_id, s.name AS service_name,
                s.category, cs.year, cs.period, cs.status, cs.payment_status,
                cs.amount_billed, cs.amount_paid, cs.amount_remaining,
                cs.due_date, cs.completed_date, cs.notes, cs.created_by
            FROM client_service cs
            JOIN service s ON cs.service_id = s.service_id
            WHERE cs.client_id = %s
            ORDER BY cs.year DESC, cs.period ASC
        """, (client_id,))
        return cursor.fetchall()
    finally:
        cursor.close()
        conn.close()


@app.post("/api/clients/{client_id}/services", status_code=201)
def add_client_service(
    client_id: int,
    service: ClientServiceCreate,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            INSERT INTO client_service
                (client_id, service_id, year, period, status,
                 amount_billed, amount_paid, due_date, notes, created_by)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """, (
            client_id, service.service_id, service.year, service.period,
            service.status, service.amount_billed, service.amount_paid,
            service.due_date, service.notes, int(current_user["sub"])
        ))
        conn.commit()
        log_action(int(current_user["sub"]), "add_service", client_id, {
            "service_id": service.service_id,
            "year": service.year,
            "period": service.period
        })
        return {"message": "Η υπηρεσία προστέθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


@app.put("/api/clients/{client_id}/services/{service_id}/status")
def update_service_status(
    client_id: int,
    service_id: int,
    year: int,
    period: str,
    status: str,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        completed_date = "CURDATE()" if status == "completed" else "NULL"
        cursor.execute(f"""
            UPDATE client_service
            SET status = %s,
                completed_date = {completed_date},
                updated_by = %s,
                updated_at = NOW()
            WHERE client_id = %s AND service_id = %s AND year = %s AND period = %s
        """, (status, int(current_user["sub"]), client_id, service_id, year, period))
        conn.commit()
        log_action(int(current_user["sub"]), "update_service", client_id, {
            "service_id": service_id,
            "new_status": status
        })
        return {"message": "Το status ενημερώθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


@app.put("/api/clients/{client_id}/services/{service_id}/payment")
def update_payment_status(
    client_id: int,
    service_id: int,
    year: int,
    period: str,
    payment_status: str,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE client_service
            SET payment_status = %s,
                updated_by = %s,
                updated_at = NOW()
            WHERE client_id = %s AND service_id = %s AND year = %s AND period = %s
        """, (payment_status, int(current_user["sub"]), client_id, service_id, year, period))
        conn.commit()
        return {"message": "Η κατάσταση πληρωμής ενημερώθηκε"}
    finally:
        cursor.close()
        conn.close()


@app.put("/api/clients/{client_id}/services/{service_id}")
def update_client_service(
    client_id: int,
    service_id: int,
    year: int,
    period: str,
    data: ClientServiceUpdate,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE client_service
            SET amount_billed = %s, amount_paid = %s, due_date = %s,
                notes = %s, updated_by = %s, updated_at = NOW()
            WHERE client_id = %s AND service_id = %s AND year = %s AND period = %s
        """, (
            data.amount_billed, data.amount_paid, data.due_date, data.notes,
            int(current_user["sub"]), client_id, service_id, year, period
        ))
        conn.commit()

        if data.amount_paid is not None and data.amount_billed is not None:
            if data.amount_billed > 0 and data.amount_paid >= data.amount_billed:
                new_payment_status = "paid"
            else:
                new_payment_status = "unpaid"
            cursor.execute("""
                UPDATE client_service
                SET payment_status = %s
                WHERE client_id = %s AND service_id = %s AND year = %s AND period = %s
            """, (new_payment_status, client_id, service_id, year, period))
            conn.commit()

        if data.due_date:
            cursor.execute("""
                UPDATE client_service
                SET status = 'pending'
                WHERE client_id = %s AND service_id = %s AND year = %s AND period = %s
                  AND status = 'overdue'
                  AND due_date > CURDATE()
            """, (client_id, service_id, year, period))
            conn.commit()

        log_action(int(current_user["sub"]), "edit_service", client_id, {
            "service_id": service_id,
            "year": year,
            "period": period
        })

        return {"message": "Η υπηρεσία ενημερώθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


@app.delete("/api/clients/{client_id}/services/{service_id}")
def delete_client_service(
    client_id: int,
    service_id: int,
    year: int,
    period: str,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor()
        cursor.execute("""
            DELETE FROM client_service
            WHERE client_id = %s AND service_id = %s AND year = %s AND period = %s
        """, (client_id, service_id, year, period))
        conn.commit()
        log_action(int(current_user["sub"]), "delete_service", client_id, {
            "service_id": service_id,
            "year": year,
            "period": period
        })
        return {"message": "Η υπηρεσία διαγράφηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


# ============================================================
# SERVICES (κατάλογος)
# ============================================================
@app.get("/api/services")
def get_services(current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT service_id, name, category, description, is_recurring
            FROM service WHERE is_active = TRUE
            ORDER BY category, name
        """)
        return cursor.fetchall()
    finally:
        cursor.close()
        conn.close()


# ============================================================
# EMPLOYEES
# ============================================================
@app.get("/api/employees")
def get_employees(current_user: dict = Depends(get_admin_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT employee_id, username, full_name, email,
                   phone, role, is_active, last_login
            FROM employee
            ORDER BY role DESC, full_name ASC
        """)
        employees = cursor.fetchall()
        for emp in employees:
            emp["access"] = ["services"]
        return employees
    finally:
        cursor.close()
        conn.close()


@app.post("/api/employees", status_code=201)
def create_employee(employee: EmployeeCreate, current_user: dict = Depends(get_admin_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        hashed = hash_password(employee.password)
        cursor.execute("""
            INSERT INTO employee (username, password_hash, full_name, email, phone, role)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (employee.username, hashed, employee.full_name,
              employee.email, employee.phone, employee.role))
        conn.commit()
        return {"message": "Ο υπάλληλος δημιουργήθηκε επιτυχώς"}
    except Exception as e:
        if "Duplicate entry" in str(e):
            raise HTTPException(status_code=409, detail="Το username υπάρχει ήδη")
        raise HTTPException(status_code=500, detail="Σφάλμα κατά τη δημιουργία υπαλλήλου")
    finally:
        cursor.close()
        conn.close()


# ============================================================
# DASHBOARD
# ============================================================
@app.get("/api/dashboard")
def get_dashboard(current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        role = current_user.get("role")

        cursor.execute("SELECT COUNT(*) as total FROM client WHERE is_active = TRUE")
        total_clients = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT COUNT(*) as total FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            WHERE cs.status IN ('pending', 'in_progress')
              AND cs.due_date BETWEEN CURDATE() AND DATE_ADD(CURDATE(), INTERVAL 7 DAY)
              AND c.is_active = TRUE
        """)
        week_pending = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT COUNT(*) as total FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            WHERE cs.status = 'overdue' AND c.is_active = TRUE
        """)
        overdue = cursor.fetchone()["total"]

        cursor.execute("""
            SELECT client_id,
                CASE
                    WHEN company_name IS NOT NULL AND company_name != '' THEN company_name
                    ELSE CONCAT(last_name, ' ', COALESCE(first_name, ''))
                END AS name,
                client_type, created_at
            FROM client
            WHERE is_active = TRUE
            ORDER BY created_at DESC
            LIMIT 5
        """)
        recent_clients = cursor.fetchall()
        for row in recent_clients:
            if row["created_at"]:
                row["created_at"] = row["created_at"].strftime("%d/%m/%Y")

        cursor.execute("""
            SELECT
                c.client_id,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    ELSE CONCAT(c.last_name, ' ', COALESCE(c.first_name, ''))
                END AS pelatis,
                s.name AS ypiresia,
                cs.due_date,
                DATEDIFF(cs.due_date, CURDATE()) AS days_left,
                cs.status
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            JOIN service s ON cs.service_id = s.service_id
            WHERE cs.status IN ('pending', 'in_progress', 'overdue')
                AND cs.payment_status = 'unpaid'
                AND cs.due_date IS NOT NULL
                AND c.is_active = TRUE
            ORDER BY cs.due_date ASC
            LIMIT 5
        """)
        urgent = cursor.fetchall()
        for row in urgent:
            if row["due_date"]:
                row["due_date"] = row["due_date"].strftime("%d/%m/%Y")

        result = {
            "total_clients": total_clients,
            "week_pending": week_pending,
            "overdue": overdue,
            "recent_clients": recent_clients,
            "urgent": urgent,
        }

        if role == "admin":
            cursor.execute("""
                SELECT COALESCE(SUM(amount_remaining), 0) as total
                FROM client_service cs
                JOIN client c ON cs.client_id = c.client_id
                WHERE cs.amount_remaining > 0
                  AND c.is_active = TRUE
            """)
            unpaid_total = cursor.fetchone()["total"]

            cursor.execute("""
                SELECT COALESCE(SUM(amount_paid), 0) as total
                FROM client_service cs
                JOIN client c ON cs.client_id = c.client_id
                WHERE MONTH(cs.created_at) = MONTH(CURDATE())
                  AND YEAR(cs.created_at) = YEAR(CURDATE())
                  AND cs.amount_paid > 0
                  AND c.is_active = TRUE
            """)
            income_this_month = cursor.fetchone()["total"]

            cursor.execute("""
                SELECT COALESCE(SUM(amount_paid), 0) as total
                FROM client_service cs
                JOIN client c ON cs.client_id = c.client_id
                WHERE MONTH(cs.created_at) = MONTH(DATE_SUB(CURDATE(), INTERVAL 1 MONTH))
                  AND YEAR(cs.created_at) = YEAR(DATE_SUB(CURDATE(), INTERVAL 1 MONTH))
                  AND cs.amount_paid > 0
                  AND c.is_active = TRUE
            """)
            income_last_month = cursor.fetchone()["total"]

            result["unpaid_total"] = float(unpaid_total)
            result["income_this_month"] = float(income_this_month)
            result["income_last_month"] = float(income_last_month)

        return result
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΑΠΛΗΡΩΤΑ ΠΟΣΑ
# ============================================================
@app.get("/api/unpaid")
def get_unpaid(current_user: dict = Depends(get_admin_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT
                c.client_id,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    ELSE CONCAT(c.last_name, ' ', COALESCE(c.first_name, ''))
                END AS pelatis,
                s.name AS ypiresia,
                cs.year,
                cs.period,
                cs.amount_billed,
                cs.amount_paid,
                cs.amount_remaining,
                cs.due_date
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            JOIN service s ON cs.service_id = s.service_id
            WHERE cs.amount_remaining > 0
              AND c.is_active = TRUE
            ORDER BY cs.due_date ASC
        """)
        results = cursor.fetchall()
        for row in results:
            if row["due_date"]:
                row["due_date"] = row["due_date"].strftime("%d/%m/%Y")
        return results
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΕΙΔΟΠΟΙΗΣΕΙΣ
# ============================================================
@app.get("/api/notifications")
def get_notifications(current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT
                s.name AS ypiresia,
                cs.due_date,
                DATEDIFF(cs.due_date, CURDATE()) AS days_left,
                COUNT(DISTINCT cs.client_id) AS client_count
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            JOIN service s ON cs.service_id = s.service_id
            WHERE cs.status IN ('pending', 'in_progress')
              AND cs.due_date IS NOT NULL
              AND cs.due_date BETWEEN CURDATE() AND DATE_ADD(CURDATE(), INTERVAL 7 DAY)
              AND c.is_active = TRUE
            GROUP BY s.service_id, s.name, cs.due_date
            ORDER BY cs.due_date ASC
        """)
        rows = cursor.fetchall()
        for row in rows:
            if row["due_date"]:
                row["due_date"] = row["due_date"].strftime("%d/%m/%Y")
        return rows
    finally:
        cursor.close()
        conn.close()


@app.get("/api/notifications/clients")
def get_clients_by_notification(
    service: str,
    due_date: str,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT DISTINCT c.client_id, c.last_name, c.first_name, c.company_name,
                    c.client_type, c.phone, c.email
            FROM client c
            JOIN client_service cs ON c.client_id = cs.client_id
            JOIN service s ON cs.service_id = s.service_id
            WHERE s.name = %s
                AND cs.due_date = STR_TO_DATE(%s, %s)
                AND cs.status IN ('pending', 'in_progress')
                AND c.is_active = TRUE
            ORDER BY c.last_name ASC
        """, (service, due_date, '%d/%m/%Y'))
        return cursor.fetchall()
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΕΚΚΡΕΜΟΤΗΤΕΣ
# ============================================================
@app.get("/api/pending")
def get_pending(
    service: Optional[str] = None,
    status: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        query = """
            SELECT
                c.client_id,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    WHEN c.last_name IS NOT NULL AND c.first_name IS NOT NULL
                        THEN CONCAT(c.last_name, ' ', c.first_name)
                    WHEN c.last_name IS NOT NULL THEN c.last_name
                    ELSE 'Άγνωστος'
                END AS pelatis,
                s.name AS ypiresia,
                s.category AS categoria,
                cs.year,
                cs.period,
                cs.status,
                cs.payment_status,
                cs.amount_remaining,
                cs.due_date,
                DATEDIFF(cs.due_date, CURDATE()) AS days_left,
                'service' AS source_type
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            JOIN service s ON cs.service_id = s.service_id
            WHERE cs.status IN ('pending', 'in_progress', 'overdue')
            AND c.is_active = TRUE
            AND s.name = COALESCE(%s, s.name)
        """
        params = [service if service else None]
        if status:
            query += " AND cs.status = %s"
            params.append(status)
        query += " ORDER BY cs.due_date ASC"
        cursor.execute(query, params)
        rows = cursor.fetchall()
        for row in rows:
            if row["due_date"]:
                row["due_date"] = row["due_date"].strftime("%d/%m/%Y")
        return rows
    finally:
        cursor.close()
        conn.close()


# ============================================================
# EXPORT ΕΚΚΡΕΜΟΤΗΤΩΝ ΣΕ EXCEL
# ============================================================
@app.get("/api/export/pending")
def export_pending_excel(current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    WHEN c.last_name IS NOT NULL AND c.first_name IS NOT NULL
                        THEN CONCAT(c.last_name, ' ', c.first_name)
                    WHEN c.last_name IS NOT NULL THEN c.last_name
                    ELSE 'Άγνωστος'
                END AS pelatis,
                s.name AS ypiresia,
                cs.year AS etos,
                cs.period AS periodo,
                cs.status AS katastasi,
                cs.payment_status AS pliromi,
                cs.amount_billed AS xreothike,
                cs.amount_paid AS plirothike,
                cs.amount_remaining AS ypoloipo,
                cs.due_date AS prothesmia
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            JOIN service s ON cs.service_id = s.service_id
            WHERE cs.status IN ('pending', 'in_progress', 'overdue')
              AND c.is_active = TRUE
            ORDER BY cs.due_date ASC
        """)
        rows = cursor.fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Εκκρεμότητες"
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_alignment = Alignment(horizontal="center")
        headers = ["Πελάτης", "Υπηρεσία", "Έτος", "Περίοδος", "Κατάσταση", "Πληρωμή",
                   "Χρεώθηκε (€)", "Πληρώθηκε (€)", "Υπόλοιπο (€)", "Προθεσμία"]
        ws.append(headers)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        status_map = {
            "pending": "Εκκρεμεί", "in_progress": "Σε εξέλιξη",
            "overdue": "Ληξιπρόθεσμο", "completed": "Ολοκληρώθηκε",
        }
        payment_map = {"unpaid": "Απλήρωτο", "paid": "Εξοφλήθηκε"}

        for row in rows:
            ws.append([
                row["pelatis"] or "Άγνωστος",
                row["ypiresia"],
                row["etos"],
                row["periodo"],
                status_map.get(row["katastasi"], row["katastasi"]),
                payment_map.get(row["pliromi"], row["pliromi"]),
                float(row["xreothike"]) if row["xreothike"] else 0,
                float(row["plirothike"]) if row["plirothike"] else 0,
                float(row["ypoloipo"]) if row["ypoloipo"] else 0,
                row["prothesmia"].strftime("%d/%m/%Y") if row["prothesmia"] else "-",
            ])
        column_widths = [25, 25, 8, 15, 15, 15, 15, 15, 15, 15]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=ekkremotites.xlsx"}
        )
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΑΝΑΦΟΡΕΣ ΕΣΟΔΩΝ ΑΝΑ ΥΠΑΛΛΗΛΟ
# ============================================================
@app.get("/api/reports/employee-revenue")
def get_employee_revenue(
    service_id: Optional[int] = None,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        query = """
            SELECT
                e.employee_id,
                e.full_name,
                YEAR(cs.created_at) AS year,
                MONTH(cs.created_at) AS month,
                COUNT(*) AS total_services,
                COALESCE(SUM(cs.amount_billed), 0) AS total_billed,
                COALESCE(SUM(CASE WHEN cs.payment_status = 'paid' THEN cs.amount_paid ELSE 0 END), 0) AS total_paid,
                COALESCE(SUM(cs.amount_remaining), 0) AS total_pending
            FROM client_service cs
            JOIN employee e ON cs.created_by = e.employee_id
            JOIN client c ON cs.client_id = c.client_id
            WHERE c.is_active = TRUE
        """
        params = []
        if service_id:
            query += " AND cs.service_id = %s"
            params.append(service_id)
        query += """
            GROUP BY e.employee_id, e.full_name, YEAR(cs.created_at), MONTH(cs.created_at)
            ORDER BY year DESC, month DESC, e.full_name ASC
        """
        cursor.execute(query, params)
        results = cursor.fetchall()
        return results
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΑΝΑΦΟΡΑ ΕΞΑΓΩΓΗΣ ΔΕΔΟΜΕΝΩΝ ΑΝΑ ΥΠΗΡΕΣΙΑ
# ============================================================
@app.get("/api/reports/service-export")
def get_service_export(
    service_id: int,
    base_year: int,
    current_year: int,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT DISTINCT
                c.client_id,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    ELSE CONCAT(c.last_name, ' ', COALESCE(c.first_name, ''))
                END AS pelatis,
                c.phone, c.afm
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            WHERE cs.service_id = %s
            AND YEAR(cs.created_at) = %s
            AND cs.status IN ('completed', 'overdue')
            AND c.is_active = TRUE
        """, (service_id, base_year))
        base_clients = cursor.fetchall()

        cursor.execute("""
            SELECT DISTINCT cs.client_id
            FROM client_service cs
            WHERE cs.service_id = %s AND YEAR(cs.created_at) = %s
        """, (service_id, current_year))
        current_year_ids = {row['client_id'] for row in cursor.fetchall()}

        done = []
        pending = []
        for client in base_clients:
            afm = decrypt(client['afm']) if client['afm'] else None
            entry = {
                'client_id': client['client_id'],
                'pelatis': client['pelatis'],
                'phone': client['phone'],
                'afm': afm,
            }
            if client['client_id'] in current_year_ids:
                done.append(entry)
            else:
                pending.append(entry)

        return {
            'base_year': base_year,
            'current_year': current_year,
            'total_base': len(base_clients),
            'total_done': len(done),
            'total_pending': len(pending),
            'done': done,
            'pending': pending,
        }
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΕΓΓΡΑΦΑ ΠΕΛΑΤΗ
# ============================================================
@app.get("/api/clients/{client_id}/documents")
def get_client_documents(client_id: int, current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT d.document_id, d.client_id, d.service_id, d.year, d.period,
                   d.filename, d.original_name, d.file_type, d.file_size,
                   d.uploaded_at, e.full_name AS uploaded_by_name,
                   s.name AS service_name
            FROM client_document d
            LEFT JOIN employee e ON d.uploaded_by = e.employee_id
            LEFT JOIN service s ON d.service_id = s.service_id
            WHERE d.client_id = %s
            ORDER BY d.uploaded_at DESC
        """, (client_id,))
        results = cursor.fetchall()
        for row in results:
            if row["uploaded_at"]:
                row["uploaded_at"] = row["uploaded_at"].strftime("%d/%m/%Y %H:%M")
        return results
    finally:
        cursor.close()
        conn.close()


@app.post("/api/clients/{client_id}/documents", status_code=201)
async def upload_document(
    client_id: int,
    file: UploadFile = File(...),
    service_id: Optional[int] = Form(None),
    year: Optional[int] = Form(None),
    period: Optional[str] = Form(None),
    current_user: dict = Depends(get_current_user)
):
    if DEMO_MODE:
        raise HTTPException(status_code=403, detail="Το ανέβασμα αρχείων είναι απενεργοποιημένο στο demo")
    allowed_types = ["application/pdf", "image/png", "image/jpeg", "image/jpg"]
    if file.content_type not in allowed_types:
        raise HTTPException(status_code=400, detail="Επιτρέπονται μόνο PDF, PNG, JPG αρχεία")

    contents = await file.read()
    file_size = len(contents)
    if file_size > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Το αρχείο δεν μπορεί να ξεπερνά τα 10MB")

    ext = os.path.splitext(file.filename)[1]
    unique_filename = f"{uuid.uuid4()}{ext}"
    file_path = os.path.join(UPLOAD_DIR, unique_filename)

    with open(file_path, "wb") as f:
        f.write(contents)

    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO client_document
                (client_id, service_id, year, period, filename, original_name,
                 file_type, file_size, uploaded_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (
            client_id, service_id, year, period,
            unique_filename, file.filename,
            file.content_type, file_size,
            int(current_user["sub"])
        ))
        conn.commit()
        return {"message": "Το αρχείο ανέβηκε επιτυχώς", "filename": unique_filename}
    finally:
        cursor.close()
        conn.close()


@app.delete("/api/clients/{client_id}/documents/{document_id}")
def delete_document(
    client_id: int,
    document_id: int,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT filename FROM client_document
            WHERE document_id = %s AND client_id = %s
        """, (document_id, client_id))
        doc = cursor.fetchone()
        if not doc:
            raise HTTPException(status_code=404, detail="Το αρχείο δεν βρέθηκε")

        file_path = os.path.join(UPLOAD_DIR, doc["filename"])
        if os.path.exists(file_path):
            os.remove(file_path)

        cursor.execute("""
            DELETE FROM client_document WHERE document_id = %s
        """, (document_id,))
        conn.commit()
        return {"message": "Το αρχείο διαγράφηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΕΠΕΞΕΡΓΑΣΙΑ ΥΠΑΛΛΗΛΟΥ
# ============================================================
def ensure_not_demo_account(cursor, employee_id: int):
    """In demo mode, keep the shared demo logins usable for every visitor."""
    if not DEMO_MODE:
        return
    cursor.execute("SELECT username FROM employee WHERE employee_id = %s", (employee_id,))
    row = cursor.fetchone()
    username = (row["username"] if isinstance(row, dict) else row[0]) if row else None
    if username in DEMO_PROTECTED_USERNAMES:
        raise HTTPException(status_code=403, detail="Οι λογαριασμοί demo δεν μπορούν να αλλάξουν")


@app.put("/api/employees/{employee_id}")
def update_employee(
    employee_id: int,
    employee: EmployeeUpdate,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor()
        ensure_not_demo_account(cursor, employee_id)
        if employee.password:
            hashed = hash_password(employee.password)
            cursor.execute("""
                UPDATE employee
                SET username = %s, password_hash = %s, full_name = %s,
                    email = %s, phone = %s, role = %s
                WHERE employee_id = %s
            """, (employee.username, hashed, employee.full_name,
                  employee.email, employee.phone, employee.role, employee_id))
        else:
            cursor.execute("""
                UPDATE employee
                SET username = %s, full_name = %s, email = %s,
                    phone = %s, role = %s
                WHERE employee_id = %s
            """, (employee.username, employee.full_name,
                  employee.email, employee.phone, employee.role, employee_id))
        conn.commit()
        return {"message": "Ο υπάλληλος ενημερώθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΑΠΕΝΕΡΓΟΠΟΙΗΣΗ ΥΠΑΛΛΗΛΟΥ
# ============================================================
@app.delete("/api/employees/{employee_id}")
def toggle_employee_status(
    employee_id: int,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        ensure_not_demo_account(cursor, employee_id)
        cursor.execute("SELECT is_active FROM employee WHERE employee_id = %s", (employee_id,))
        emp = cursor.fetchone()
        if not emp:
            raise HTTPException(status_code=404, detail="Ο υπάλληλος δεν βρέθηκε")
        new_status = not emp["is_active"]
        cursor.execute("UPDATE employee SET is_active = %s WHERE employee_id = %s", (new_status, employee_id))
        conn.commit()
        msg = "ενεργοποιήθηκε" if new_status else "απενεργοποιήθηκε"
        return {"message": f"Ο υπάλληλος {msg} επιτυχώς", "is_active": new_status}
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΙΣΤΟΡΙΚΟ ΚΙΝΗΣΕΩΝ ΥΠΑΛΛΗΛΟΥ
# ============================================================
@app.get("/api/employees/{employee_id}/actions")
def get_employee_actions(
    employee_id: int,
    action_type: Optional[str] = None,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης με τη βάση")
    try:
        cursor = conn.cursor(dictionary=True)
        query = """
            SELECT
                a.employee_id,
                a.timestamp,
                a.action_type,
                a.client_id,
                a.details,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    WHEN c.last_name IS NOT NULL AND c.first_name IS NOT NULL
                        THEN CONCAT(c.last_name, ' ', c.first_name)
                    WHEN c.last_name IS NOT NULL THEN c.last_name
                    ELSE NULL
                END AS client_name
            FROM action a
            LEFT JOIN client c ON a.client_id = c.client_id
            WHERE a.employee_id = %s
        """
        params = [employee_id]
        if action_type and action_type != "all":
            query += " AND a.action_type = %s"
            params.append(action_type)
        query += " ORDER BY a.timestamp DESC LIMIT 100"
        cursor.execute(query, params)
        results = cursor.fetchall()
        for row in results:
            if row["timestamp"]:
                row["timestamp"] = row["timestamp"].strftime("%d/%m/%Y %H:%M")
            if row["details"] and isinstance(row["details"], str):
                try:
                    row["details"] = json.loads(row["details"])
                except:
                    pass
        return results
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΠΛΗΡΩΜΕΣ ΕΠΙΧΕΙΡΗΣΕΩΝ
# ============================================================
class PaymentCreate(BaseModel):
    amount: float
    payment_date: str
    notes: Optional[str] = None


@app.get("/api/clients/{client_id}/payments")
def get_client_payments(client_id: int, year: int, current_user: dict = Depends(get_current_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT p.payment_id, p.amount, p.payment_date,
                   p.notes, p.created_at, e.full_name AS created_by_name
            FROM client_payment p
            JOIN employee e ON p.created_by = e.employee_id
            WHERE p.client_id = %s AND p.year = %s
            ORDER BY p.payment_date ASC
        """, (client_id, year))
        rows = cursor.fetchall()
        for row in rows:
            if row["payment_date"]:
                row["payment_date"] = row["payment_date"].strftime("%d/%m/%Y")
            if row["created_at"]:
                row["created_at"] = row["created_at"].strftime("%d/%m/%Y %H:%M")
        return rows
    finally:
        cursor.close()
        conn.close()


@app.post("/api/clients/{client_id}/payments", status_code=201)
def add_client_payment(
    client_id: int,
    year: int,
    service_id: int,
    data: PaymentCreate,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            INSERT INTO client_payment
                (client_id, service_id, year, amount, payment_date, notes, created_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (client_id, service_id, year, data.amount,
              data.payment_date, data.notes, int(current_user["sub"])))

        cursor.execute("""
            UPDATE client_service
            SET amount_paid = (
                SELECT COALESCE(SUM(amount), 0)
                FROM client_payment
                WHERE client_id = %s AND service_id = %s AND year = %s
            )
            WHERE client_id = %s AND service_id = %s AND year = %s
        """, (client_id, service_id, year, client_id, service_id, year))

        cursor.execute("""
            UPDATE client_service
            SET payment_status = CASE
                WHEN amount_paid >= amount_billed THEN 'paid'
                ELSE 'unpaid'
            END
            WHERE client_id = %s AND service_id = %s AND year = %s
        """, (client_id, service_id, year))

        conn.commit()
        log_action(int(current_user["sub"]), "edit_service", client_id, {
            "service_id": service_id, "year": year, "payment": data.amount
        })
        return {"message": "Η δόση καταχωρήθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


@app.delete("/api/clients/{client_id}/payments/{payment_id}")
def delete_client_payment(
    client_id: int,
    payment_id: int,
    year: int,
    service_id: int,
    current_user: dict = Depends(get_current_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor()
        cursor.execute("""
            DELETE FROM client_payment
            WHERE payment_id = %s AND client_id = %s
        """, (payment_id, client_id))

        cursor.execute("""
            UPDATE client_service
            SET amount_paid = (
                SELECT COALESCE(SUM(amount), 0)
                FROM client_payment
                WHERE client_id = %s AND service_id = %s AND year = %s
            )
            WHERE client_id = %s AND service_id = %s AND year = %s
        """, (client_id, service_id, year, client_id, service_id, year))

        cursor.execute("""
            UPDATE client_service
            SET payment_status = CASE
                WHEN amount_paid >= amount_billed THEN 'paid'
                ELSE 'unpaid'
            END
            WHERE client_id = %s AND service_id = %s AND year = %s
        """, (client_id, service_id, year))

        conn.commit()
        return {"message": "Η δόση διαγράφηκε"}
    finally:
        cursor.close()
        conn.close()


# ============================================================
# ΔΙΑΦΟΡΑ ΕΣΟΔΑ
# ============================================================
class MiscIncomeCreate(BaseModel):
    amount: float
    employee_name: str
    description: Optional[str] = None
    income_date: str


@app.get("/api/misc-income")
def get_misc_income(
    year: Optional[int] = None,
    month: Optional[int] = None,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor(dictionary=True)
        query = """
            SELECT income_id, amount, employee_name, description,
                   income_date, created_at
            FROM misc_income
            WHERE 1=1
        """
        params = []
        if year:
            query += " AND YEAR(income_date) = %s"
            params.append(year)
        if month:
            query += " AND MONTH(income_date) = %s"
            params.append(month)
        query += " ORDER BY income_date DESC"
        cursor.execute(query, params)
        rows = cursor.fetchall()
        for row in rows:
            if row["income_date"]:
                row["income_date"] = row["income_date"].strftime("%d/%m/%Y")
            if row["created_at"]:
                row["created_at"] = row["created_at"].strftime("%d/%m/%Y %H:%M")
        return rows
    finally:
        cursor.close()
        conn.close()


@app.post("/api/misc-income", status_code=201)
def add_misc_income(
    data: MiscIncomeCreate,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO misc_income
                (amount, employee_name, description, income_date, created_by)
            VALUES (%s, %s, %s, %s, %s)
        """, (data.amount, data.employee_name, data.description,
              data.income_date, int(current_user["sub"])))
        conn.commit()
        return {"message": "Καταχωρήθηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()


@app.delete("/api/misc-income/{income_id}")
def delete_misc_income(
    income_id: int,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM misc_income WHERE income_id = %s", (income_id,))
        conn.commit()
        return {"message": "Διαγράφηκε επιτυχώς"}
    finally:
        cursor.close()
        conn.close()

@app.get("/api/reports/payments-by-year")
def get_payments_by_year(current_user: dict = Depends(get_admin_user)):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT
                cs.year,
                c.client_id,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    ELSE CONCAT(c.last_name, ' ', COALESCE(c.first_name, ''))
                END AS pelatis,
                cs.amount_billed,
                cs.amount_paid,
                cs.amount_remaining,
                cs.payment_status
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            WHERE cs.service_id = 2
              AND c.is_active = TRUE
            ORDER BY cs.year DESC, c.last_name ASC
        """)
        rows = cursor.fetchall()
        return rows
    finally:
        cursor.close()
        conn.close()


@app.get("/api/reports/declarations-unpaid")
def get_declarations_unpaid(
    year: int,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT
                c.client_id,
                CASE
                    WHEN c.company_name IS NOT NULL AND c.company_name != '' THEN c.company_name
                    ELSE CONCAT(c.last_name, ' ', COALESCE(c.first_name, ''))
                END AS pelatis,
                cs.amount_billed,
                cs.amount_paid,
                cs.amount_remaining,
                cs.payment_status,
                MONTH(cs.created_at) AS month
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            WHERE cs.service_id = 1
              AND YEAR(cs.created_at) = %s
              AND cs.amount_remaining > 0
              AND c.is_active = TRUE
            ORDER BY c.last_name ASC
        """, (year,))
        return cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

@app.get("/api/reports/annual-summary")
def get_annual_summary(
    year: int,
    current_user: dict = Depends(get_admin_user)
):
    conn = get_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Σφάλμα σύνδεσης")
    try:
        cursor = conn.cursor(dictionary=True)

        # Φορολογικές
        cursor.execute("""
            SELECT COALESCE(SUM(amount_paid), 0) as total
            FROM client_service cs
            JOIN client c ON cs.client_id = c.client_id
            WHERE cs.service_id = 1
              AND YEAR(cs.created_at) = %s
              AND c.is_active = TRUE
        """, (year,))
        declarations = float(cursor.fetchone()["total"])

        # Πληρωμές Επιχειρήσεων (βάσει payment_date)
        cursor.execute("""
            SELECT COALESCE(SUM(amount), 0) as total
            FROM client_payment
            WHERE service_id = 2
              AND YEAR(payment_date) = %s
        """, (year,))
        payments = float(cursor.fetchone()["total"])

        # Διάφορα Έσοδα
        cursor.execute("""
            SELECT COALESCE(SUM(amount), 0) as total
            FROM misc_income
            WHERE YEAR(income_date) = %s
        """, (year,))
        misc = float(cursor.fetchone()["total"])

        return {
            "year": year,
            "declarations": declarations,
            "payments": payments,
            "misc": misc,
            "total": declarations + payments + misc
        }
    finally:
        cursor.close()
        conn.close()


# ============================================================
# LIVE DEMO: ΚΑΘΗΜΕΡΙΝΗ ΕΠΑΝΑΦΟΡΑ ΔΕΔΟΜΕΝΩΝ
# ============================================================
@app.get("/api/cron/reset-demo")
def reset_demo(authorization: Optional[str] = Header(None)):
    """Called daily by Vercel Cron. Rebuilds the database with fake data."""
    cron_secret = os.getenv("CRON_SECRET")
    if not DEMO_MODE or not cron_secret:
        raise HTTPException(status_code=404, detail="Not found")
    if not authorization or not secrets.compare_digest(authorization, f"Bearer {cron_secret}"):
        raise HTTPException(status_code=401, detail="Unauthorized")
    from db.seed import reset_database
    counts = reset_database()
    return {"message": "Demo database reset", "rows": counts}


@app.get("/api/health")
def health():
    return {"status": "ok"}
